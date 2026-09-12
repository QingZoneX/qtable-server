"""Generate real Office Open XML (.xlsx) exports for SmartTable data."""
from __future__ import annotations

import json
from datetime import datetime, timezone
from io import BytesIO
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill


XLSX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _requested_items(
    items: Sequence[Dict[str, Any]],
    ids: Optional[Sequence[str]],
    *,
    label: str,
) -> List[Dict[str, Any]]:
    if ids is None:
        return list(items)

    by_id = {
        str(item.get("id")): item
        for item in items
        if item.get("id") is not None
    }
    requested: List[Dict[str, Any]] = []
    missing: List[str] = []
    for item_id in ids:
        key = str(item_id)
        item = by_id.get(key)
        if item is None:
            missing.append(key)
        else:
            requested.append(item)
    if missing:
        raise ValueError(f"Unknown {label} ids: {', '.join(missing[:10])}")
    return requested


def _as_text(value: Any, options: Optional[Dict[str, str]] = None) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        return ", ".join(
            part
            for part in (_as_text(item, options) for item in value)
            if part != ""
        )
    if isinstance(value, dict):
        for key in ("label", "name", "title"):
            candidate = value.get(key)
            if candidate is not None and candidate != "":
                return str(candidate)
        item_id = value.get("id")
        if item_id is not None:
            raw_id = str(item_id)
            return (options or {}).get(raw_id, raw_id)
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    raw = str(value)
    if options:
        return options.get(raw, raw)
    return raw


def _as_number(value: Any) -> Optional[int | float]:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return value
    if isinstance(value, str):
        cleaned = value.strip().replace(",", "")
        if not cleaned:
            return None
        try:
            number = float(cleaned)
        except ValueError:
            return None
        return int(number) if number.is_integer() else number
    return None


def _as_datetime(value: Any) -> Optional[datetime]:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, (int, float)):
        seconds = float(value)
        if abs(seconds) > 100_000_000_000:
            seconds /= 1000.0
        try:
            parsed = datetime.fromtimestamp(seconds, tz=timezone.utc)
        except (OverflowError, OSError, ValueError):
            return None
    elif isinstance(value, str):
        cleaned = value.strip()
        if not cleaned:
            return None
        try:
            numeric = float(cleaned)
        except ValueError:
            numeric = None
        if numeric is not None:
            return _as_datetime(numeric)
        normalized = cleaned[:-1] + "+00:00" if cleaned.endswith("Z") else cleaned
        try:
            parsed = datetime.fromisoformat(normalized)
        except ValueError:
            return None
    else:
        return None

    if parsed.tzinfo is not None:
        parsed = parsed.astimezone(timezone.utc).replace(tzinfo=None)
    return parsed


def _excel_date_format(field: Dict[str, Any]) -> str:
    configured = str((field.get("property") or {}).get("format") or "").strip()
    if not configured:
        return "yyyy-mm-dd"
    # QTable/Dayjs tokens are close enough to Excel tokens for the supported
    # date formats. Excel uses lowercase yyyy/dd/hh but is case-insensitive for
    # several tokens; normalize the common forms for predictable output.
    replacements = (
        ("YYYY", "yyyy"),
        ("YY", "yy"),
        ("DD", "dd"),
        ("HH", "hh"),
    )
    result = configured
    for source, target in replacements:
        result = result.replace(source, target)
    return result


def _cell_payload(field: Dict[str, Any], raw_value: Any) -> Tuple[Any, Optional[str], bool]:
    """Return (value, number_format, force_text)."""

    field_type = str(field.get("type") or "text")
    options = {
        str(option.get("id")): str(option.get("label") or option.get("id") or "")
        for option in (field.get("options") or [])
        if isinstance(option, dict) and option.get("id") is not None
    }

    if field_type == "date":
        parsed = _as_datetime(raw_value)
        if parsed is not None:
            return parsed, _excel_date_format(field), False
        return _as_text(raw_value, options), None, True

    if field_type in {"number", "progress", "rating"}:
        number = _as_number(raw_value)
        if number is not None:
            return number, None, False
        return _as_text(raw_value, options), None, True

    if field_type == "autoNumber":
        number = _as_number(raw_value)
        if number is not None:
            digits = int((field.get("property") or {}).get("digits") or 0)
            return number, ("0" * digits) if digits > 0 else None, False
        return _as_text(raw_value, options), None, True

    # Formula fields are materialized by the store. Preserve scalar result
    # types rather than converting every formula result to display text.
    if field_type == "formula":
        if isinstance(raw_value, bool):
            return raw_value, None, False
        if isinstance(raw_value, (int, float)) and not isinstance(raw_value, bool):
            return raw_value, None, False
        if isinstance(raw_value, datetime):
            value = raw_value.replace(tzinfo=None) if raw_value.tzinfo else raw_value
            return value, "yyyy-mm-dd hh:mm:ss", False
        return _as_text(raw_value, options), None, True

    # Text-like fields intentionally stay strings. Explicitly forcing the cell
    # data type to string also prevents values beginning with '=' from becoming
    # executable Excel formulas (formula injection).
    return _as_text(raw_value, options), None, True


def build_xlsx_export(
    fields: Sequence[Dict[str, Any]],
    records: Sequence[Dict[str, Any]],
    *,
    field_ids: Optional[Sequence[str]] = None,
    record_ids: Optional[Sequence[str]] = None,
) -> bytes:
    """Build a real .xlsx workbook while preserving meaningful cell types."""

    selected_fields = _requested_items(fields, field_ids, label="field")
    selected_records = _requested_items(records, record_ids, label="record")

    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("QTable")
    sheet.freeze_panes = "A2"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(fill_type="solid", fgColor="3B82F6")
    header_alignment = Alignment(vertical="center")

    header_cells: List[WriteOnlyCell] = []
    for field in selected_fields:
        cell = WriteOnlyCell(sheet, value=str(field.get("name") or field.get("id") or ""))
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.data_type = "s"
        header_cells.append(cell)
    sheet.append(header_cells)

    for record in selected_records:
        row: List[WriteOnlyCell] = []
        for field in selected_fields:
            field_id = str(field.get("id"))
            value, number_format, force_text = _cell_payload(field, record.get(field_id))
            cell = WriteOnlyCell(sheet, value=value)
            if force_text:
                cell.data_type = "s"
            if number_format:
                cell.number_format = number_format
            row.append(cell)
        sheet.append(row)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
