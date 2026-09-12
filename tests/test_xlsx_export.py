from __future__ import annotations

from datetime import datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.services.xlsx_export import build_xlsx_export


FIELDS = [
    {"id": "title", "name": "任务", "type": "text"},
    {"id": "amount", "name": "金额", "type": "number"},
    {
        "id": "status",
        "name": "状态",
        "type": "select",
        "options": [
            {"id": "todo", "label": "待处理"},
            {"id": "done", "label": "已完成"},
        ],
    },
    {
        "id": "tags",
        "name": "标签",
        "type": "multiSelect",
        "options": [
            {"id": "a", "label": "A"},
            {"id": "b", "label": "B"},
        ],
    },
    {
        "id": "start",
        "name": "开始时间",
        "type": "date",
        "property": {"format": "YYYY-MM-DD"},
    },
    {"id": "phone", "name": "电话", "type": "phone"},
    {
        "id": "seq",
        "name": "编号",
        "type": "autoNumber",
        "property": {"digits": 4},
    },
    {"id": "formula", "name": "公式结果", "type": "formula"},
]


RECORDS = [
    {
        "id": "r1",
        "title": "=SUM(1,2)",
        "amount": "1234.50",
        "status": "done",
        "tags": ["a", "b"],
        "start": 1785801600000,
        "phone": "0013800138000",
        "seq": 7,
        "formula": 2469,
    },
    {
        "id": "r2",
        "title": "普通任务",
        "amount": 42,
        "status": "todo",
        "tags": [],
        "start": "2026-08-08T09:30:00",
        "phone": "010-12345678",
        "seq": 8,
        "formula": "OK",
    },
]


def workbook_from_export(**kwargs):
    payload = build_xlsx_export(FIELDS, RECORDS, **kwargs)
    assert payload[:2] == b"PK"
    return load_workbook(BytesIO(payload), data_only=False)


def test_xlsx_export_is_real_workbook_and_preserves_types():
    workbook = workbook_from_export()
    sheet = workbook["QTable"]

    assert [cell.value for cell in sheet[1]] == [
        "任务",
        "金额",
        "状态",
        "标签",
        "开始时间",
        "电话",
        "编号",
        "公式结果",
    ]

    # Text beginning with '=' must stay literal text rather than an Excel formula.
    assert sheet["A2"].value == "=SUM(1,2)"
    assert sheet["A2"].data_type == "s"

    assert sheet["B2"].value == 1234.5
    assert sheet["B2"].data_type == "n"

    assert sheet["C2"].value == "已完成"
    assert sheet["D2"].value == "A, B"

    assert isinstance(sheet["E2"].value, datetime)
    assert sheet["E2"].number_format == "yyyy-MM-dd"

    # Phone numbers remain text so leading zeroes survive.
    assert sheet["F2"].value == "0013800138000"
    assert sheet["F2"].data_type == "s"

    assert sheet["G2"].value == 7
    assert sheet["G2"].number_format == "0000"

    assert sheet["H2"].value == 2469
    assert sheet["H2"].data_type == "n"


def test_xlsx_export_respects_requested_field_and_record_order():
    workbook = workbook_from_export(
        field_ids=["status", "title", "amount"],
        record_ids=["r2", "r1"],
    )
    sheet = workbook["QTable"]

    assert [cell.value for cell in sheet[1]] == ["状态", "任务", "金额"]
    assert [cell.value for cell in sheet[2]] == ["待处理", "普通任务", 42]
    assert [cell.value for cell in sheet[3]] == ["已完成", "=SUM(1,2)", 1234.5]


def test_xlsx_export_rejects_unknown_field_or_record_ids():
    with pytest.raises(ValueError, match="Unknown field ids"):
        build_xlsx_export(FIELDS, RECORDS, field_ids=["missing"])

    with pytest.raises(ValueError, match="Unknown record ids"):
        build_xlsx_export(FIELDS, RECORDS, record_ids=["missing"])


def test_xlsx_export_handles_invalid_number_or_date_as_text():
    records = [
        {
            "id": "r1",
            "amount": "not-number",
            "start": "not-date",
        }
    ]
    fields = [
        {"id": "amount", "name": "金额", "type": "number"},
        {"id": "start", "name": "开始", "type": "date"},
    ]

    payload = build_xlsx_export(fields, records)
    workbook = load_workbook(BytesIO(payload), data_only=False)
    sheet = workbook["QTable"]

    assert sheet["A2"].value == "not-number"
    assert sheet["A2"].data_type == "s"
    assert sheet["B2"].value == "not-date"
    assert sheet["B2"].data_type == "s"
